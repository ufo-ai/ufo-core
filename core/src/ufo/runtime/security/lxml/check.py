import ctypes
import gzip
import hashlib
import io
import json
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree, objectify  # type: ignore[import-untyped]

assert etree.LXML_VERSION[:3] == (6, 1, 2)
assert etree.LIBXML_VERSION == (2, 14, 6)
assert etree.LIBXSLT_VERSION == (1, 1, 43)
assert "zlib" in etree.LIBXML_FEATURES
hashes = {}
for module in (etree, objectify):
    name = module.__name__.split(".")[-1]
    path = Path(module.__file__)
    hashes[path.name] = hashlib.sha256(path.read_bytes()).hexdigest()
    if Path("/checks").exists():
        symbols = subprocess.check_output(["nm", "-D", "--defined-only", str(path)], text=True)
        assert re.search(r"\bgzwrite\b", symbols)
        assert ctypes.CDLL(f"/checks/state-{name}.so").check_zlib_state() == 0
        needed = subprocess.check_output(["readelf", "-d", str(path)], text=True)
        assert "libz.so" not in needed
receipt = Path("/usr/local/share/ufo-lxml-zlib/module-sha256.json")
if receipt.exists():
    assert hashes == json.loads(receipt.read_text())
elif Path("/out").exists():
    Path("/out/module-sha256.json").write_text(json.dumps(hashes, indent=2) + "\n")

tree = etree.ElementTree(etree.XML(b"<root><item>one</item></root>"))
assert tree.xpath("string(/root/item)") == "one"
style = etree.XML(
    b'<xsl:stylesheet version="1.0" xmlns:xsl="http://www.w3.org/1999/XSL/Transform">'
    b'<xsl:template match="/"><out><xsl:value-of select="root/item"/></out>'
    b"</xsl:template></xsl:stylesheet>"
)
assert etree.XSLT(style)(tree).getroot().text == "one"
with tempfile.TemporaryDirectory() as directory:
    root = Path(directory)
    compressed = root / "tree.xml.gz"
    tree.write(str(compressed), compression=6)
    assert etree.fromstring(gzip.decompress(compressed.read_bytes())).tag == "root"
    assert etree.parse(str(compressed), etree.XMLParser(decompress=True)).getroot().tag == "root"
    canonical = root / "canonical.xml.gz"
    tree.write_c14n(str(canonical), compression=6)
    assert gzip.decompress(canonical.read_bytes()) == b"<root><item>one</item></root>"
    stream = io.BytesIO()
    with etree.xmlfile(stream, compression=6) as writer:
        writer.write(tree.getroot())
    assert etree.fromstring(gzip.decompress(stream.getvalue())).tag == "root"
    document = Document()
    document.add_paragraph("Checked document")
    document.save(str(root / "roundtrip.docx"))
    assert Document(str(root / "roundtrip.docx")).paragraphs[0].text == "Checked document"
    for name in sys.argv[1:]:
        with zipfile.ZipFile(name) as archive:
            xml_parts = [part for part in archive.namelist() if part.endswith(".xml")]
            assert xml_parts
            for part in xml_parts:
                assert etree.fromstring(archive.read(part)) is not None
print("Both patched embedded zlib copies, XML, XPath, XSLT, gzip, and DOCX passed")
