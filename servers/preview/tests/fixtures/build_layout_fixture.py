from datetime import UTC, datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

OUTPUT = Path(__file__).with_name("fixture-layout.docx")
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}
FIXED_ZIP_TIME = (2026, 8, 24, 12, 0, 0)


def child(parent, tag):
    found = parent.find(qn(tag))
    if found is None:
        found = OxmlElement(tag)
        parent.append(found)
    return found


def width(parent, tag, value):
    node = child(parent, tag)
    node.set(qn("w:type"), "dxa")
    node.set(qn("w:w"), str(value))


def table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    properties = table._tbl.tblPr
    width(properties, "w:tblW", CONTENT_WIDTH_DXA)
    indent = child(properties, "w:tblInd")
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    layout = child(properties, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for node in list(grid):
        grid.remove(node)
    for value in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)
    for row in table.rows:
        row.height = None
        for index, cell in enumerate(row.cells):
            cell.width = Twips(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            width(properties, "w:tcW", widths[index])
            margins = child(properties, "w:tcMar")
            for side, value in CELL_MARGINS_DXA.items():
                margin = child(margins, f"w:{side}")
                margin.set(qn("w:w"), str(value))
                margin.set(qn("w:type"), "dxa")


def shade(cell, fill):
    node = child(cell._tc.get_or_add_tcPr(), "w:shd")
    node.set(qn("w:fill"), fill)


document = Document()
section = document.sections[0]
section.start_type = WD_SECTION.NEW_PAGE
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.right_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

normal = document.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
normal.paragraph_format.line_spacing = 1.1

heading = document.styles["Heading 1"]
heading.font.name = "Calibri"
heading.font.size = Pt(16)
heading.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
heading.paragraph_format.space_before = Pt(16)
heading.paragraph_format.space_after = Pt(8)

title = document.add_paragraph()
title.paragraph_format.space_after = Pt(8)
run = title.add_run("DOCX clipping and overlap layout page one")
run.bold = True
run.font.name = "Calibri"
run.font.size = Pt(20)
run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)
document.add_paragraph(
    "This fixture keeps descenders (g, j, p, q, y), ascenders, and wrapped lines visible without "
    "clipping or overlap."
)

table = document.add_table(rows=1, cols=2)
table.style = "Table Grid"
table.rows[0].cells[0].text = "Layout case"
table.rows[0].cells[1].text = "Visible content"
for cell in table.rows[0].cells:
    shade(cell, "F2F4F7")
    for cell_run in cell.paragraphs[0].runs:
        cell_run.bold = True
for label, value in (
    (
        "Wrapped cell",
        "A deliberately long line wraps across several lines. Every descender and final word must "
        "remain visible, and the row must grow rather than clip.",
    ),
    (
        "Adjacent text",
        "Line one sits above line two with normal spacing.\nLine two must not overlap line one.",
    ),
    (
        "Boundary text",
        "Left and right cell padding remain visible at the print boundary.",
    ),
):
    cells = table.add_row().cells
    cells[0].text = label
    cells[1].text = value
table_geometry(table, [2700, 6660])

document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
document.add_heading("DOCX layout page two", level=1)
document.add_paragraph(
    "A non-first-page read must return this visible text and the rendered page image."
)
second = document.add_table(rows=1, cols=3)
second.style = "Table Grid"
second.rows[0].cells[0].text = "Check"
second.rows[0].cells[1].text = "Expected"
second.rows[0].cells[2].text = "Evidence"
for cell in second.rows[0].cells:
    shade(cell, "E8EEF5")
    for cell_run in cell.paragraphs[0].runs:
        cell_run.bold = True
for values in (
    ("Page range", "Page two only", "start_page=2"),
    ("Text", "Visible", "DOCX layout page two"),
    ("Image", "PNG", "No clipped rows"),
):
    cells = second.add_row().cells
    for cell, value in zip(cells, values, strict=True):
        cell.text = value
table_geometry(second, [1800, 3780, 3780])

document.core_properties.author = "ufo"
document.core_properties.last_modified_by = "ufo"
document.core_properties.created = datetime(2026, 8, 24, 12, tzinfo=UTC)
document.core_properties.modified = datetime(2026, 8, 24, 12, tzinfo=UTC)
document.save(OUTPUT)

normalized = OUTPUT.with_suffix(".normalized.docx")
with ZipFile(OUTPUT) as source, ZipFile(normalized, "w", ZIP_DEFLATED) as target:
    for name in sorted(source.namelist()):
        info = ZipInfo(name, FIXED_ZIP_TIME)
        info.compress_type = ZIP_DEFLATED
        info.external_attr = 0o600 << 16
        target.writestr(info, source.read(name))
normalized.replace(OUTPUT)
